"""
Выгрузка результатов:
  • save_excel  — аккуратный .xlsx: шапка, фильтры, форматы цен, ссылки, сводка по категориям;
  • save_gsheet — запись в Google Таблицу через сервисный аккаунт (см. README, раздел «Google Таблицы»).
"""

import logging
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from scraper import Book

log = logging.getLogger("parser")

# Колонки итоговой таблицы: (заголовок, ширина, как достать значение из книги)
COLUMNS = [
    ("Название", 48, lambda b, rate: b.title),
    ("Категория", 18, lambda b, rate: b.category),
    ("Цена, £", 10, lambda b, rate: b.price_gbp),
    ("Цена, ₽", 11, lambda b, rate: round(b.price_gbp * rate, 2) if rate else None),
    ("Наличие", 12, lambda b, rate: "В наличии" if b.in_stock else "Нет"),
    ("Кол-во", 9, lambda b, rate: b.quantity),
    ("Рейтинг", 9, lambda b, rate: b.rating),
    ("UPC", 18, lambda b, rate: b.upc),
    ("Ссылка", 50, lambda b, rate: b.url),
    ("Обложка", 50, lambda b, rate: b.image_url),
    ("Описание", 80, lambda b, rate: b.description),
]

HEADER_FILL = PatternFill("solid", fgColor="1D4ED8")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _active_columns(books: list[Book]) -> list[tuple]:
    """Колонки, которые есть смысл показывать: без --details не выводим пустые Кол-во/UPC/Описание."""
    with_details = any(b.quantity is not None for b in books)
    detail_only = {"Кол-во", "UPC", "Описание"}
    return [c for c in COLUMNS if with_details or c[0] not in detail_only]


def build_rows(books: list[Book], rate: float | None) -> tuple[list[str], list[list]]:
    """Заголовки и строки — общие для Excel и Google Таблицы."""
    cols = _active_columns(books)
    header = [c[0] for c in cols]
    rows = [[getter(b, rate) for _, _, getter in cols] for b in books]
    return header, rows


def summarize(books: list[Book]) -> list[list]:
    """Сводка по категориям: количество, средняя/мин/макс цена, средний рейтинг, сколько в наличии."""
    groups: dict[str, list[Book]] = defaultdict(list)
    for b in books:
        groups[b.category].append(b)

    summary = []
    for category, items in sorted(groups.items(), key=lambda kv: -len(kv[1])):
        prices = [b.price_gbp for b in items]
        summary.append([
            category,
            len(items),
            round(sum(prices) / len(prices), 2),
            min(prices),
            max(prices),
            round(sum(b.rating for b in items) / len(items), 2),
            sum(b.in_stock for b in items),
        ])
    return summary


# ====================================================================== Excel

def save_excel(books: list[Book], path: Path, rate: float | None, run_info: dict) -> Path:
    wb = Workbook()

    # ---------- Лист 1: все книги ----------
    ws = wb.active
    ws.title = "Книги"
    cols = _active_columns(books)
    header, rows = build_rows(books, rate)
    ws.append(header)
    for row in rows:
        ws.append(row)

    _style_header(ws, len(header))
    for i, (name, width, _) in enumerate(cols, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = width
        if name == "Цена, £":
            _format_column(ws, i, '£#,##0.00')
        elif name == "Цена, ₽":
            _format_column(ws, i, '#,##0.00 "₽"')
        elif name in ("Рейтинг", "Кол-во", "Наличие"):
            _align_column(ws, i, "center")
        elif name in ("Ссылка", "Обложка"):
            _make_links(ws, i)
        elif name == "Описание":
            _align_column(ws, i, "left", wrap=False)

    last_row = ws.max_row
    rating_col = get_column_letter(header.index("Рейтинг") + 1)
    # Цветовая шкала рейтинга: 1 — красный, 3 — жёлтый, 5 — зелёный
    ws.conditional_formatting.add(
        f"{rating_col}2:{rating_col}{last_row}",
        ColorScaleRule(start_type="num", start_value=1, start_color="F8696B",
                       mid_type="num", mid_value=3, mid_color="FFEB84",
                       end_type="num", end_value=5, end_color="63BE7B"),
    )

    # ---------- Лист 2: сводка по категориям ----------
    ws2 = wb.create_sheet("Сводка")
    ws2.append(["Категория", "Книг", "Средняя цена, £", "Мин. цена, £", "Макс. цена, £", "Средний рейтинг", "В наличии"])
    for row in summarize(books):
        ws2.append(row)
    _style_header(ws2, 7)
    for i, width in enumerate([24, 8, 16, 13, 13, 16, 11], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = width
    for i in (3, 4, 5):
        _format_column(ws2, i, '£#,##0.00')
    _format_column(ws2, 6, '0.00')

    # ---------- Лист 3: параметры запуска ----------
    ws3 = wb.create_sheet("Параметры")
    for key, value in run_info.items():
        ws3.append([key, value])
    ws3.column_dimensions["A"].width = 28
    ws3.column_dimensions["B"].width = 60
    for cell in ws3["A"]:
        cell.font = Font(bold=True)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _style_header(ws, n_cols: int) -> None:
    """Синяя шапка, закреплённая первая строка, автофильтр."""
    for cell in ws[1][:n_cols]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(n_cols)}{ws.max_row}"


def _format_column(ws, col: int, number_format: str) -> None:
    for (cell,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
        cell.number_format = number_format


def _align_column(ws, col: int, horizontal: str, wrap: bool = False) -> None:
    for (cell,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
        cell.alignment = Alignment(horizontal=horizontal, wrap_text=wrap)


def _make_links(ws, col: int) -> None:
    """Делает адреса кликабельными."""
    for (cell,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
        if cell.value:
            cell.hyperlink = cell.value
            cell.style = "Hyperlink"


# ====================================================================== Google Таблицы

def save_gsheet(books: list[Book], rate: float | None, credentials_file: str, sheet_id: str) -> str:
    """
    Записывает книги и сводку в Google Таблицу. Старые данные на листах заменяются.
    Нужен сервисный аккаунт Google и доступ «Редактор» к таблице для его e-mail (см. README).
    """
    import gspread  # импорт здесь: без флага --gsheet библиотека не нужна

    client = gspread.service_account(filename=credentials_file)
    spreadsheet = client.open_by_key(sheet_id)

    header, rows = build_rows(books, rate)
    _write_sheet(spreadsheet, "Книги", [header] + rows)

    summary_header = ["Категория", "Книг", "Средняя цена, £", "Мин. цена, £", "Макс. цена, £", "Средний рейтинг", "В наличии"]
    _write_sheet(spreadsheet, "Сводка", [summary_header] + summarize(books))

    # Отметка о последнем обновлении — видно, что данные свежие
    _write_sheet(spreadsheet, "Параметры", [["Обновлено", datetime.now().strftime("%d.%m.%Y %H:%M")],
                                            ["Книг", len(books)]])
    return spreadsheet.url


def _write_sheet(spreadsheet, title: str, values: list[list]) -> None:
    import gspread

    try:
        ws = spreadsheet.worksheet(title)
        ws.clear()
    except gspread.WorksheetNotFound:
        ws = spreadsheet.add_worksheet(title=title, rows=len(values) + 10, cols=len(values[0]) + 2)

    # None → пустая строка: Google API не принимает null
    clean = [["" if v is None else v for v in row] for row in values]
    # RAW — значения пишутся как есть: текст вида «=...» из названия не превратится в формулу
    ws.update(clean, "A1", value_input_option="RAW")
    ws.freeze(rows=1)
    ws.format(f"A1:{get_column_letter(len(values[0]))}1",
              {"textFormat": {"bold": True, "foregroundColor": {"red": 1, "green": 1, "blue": 1}},
               "backgroundColor": {"red": 0.11, "green": 0.3, "blue": 0.85}})
